#!/usr/bin/env python3
"""Local web server for the coronary branch labeler.

Run with: python server.py
Then open http://127.0.0.1:8765
"""
import base64
import csv
import importlib.util
import io
import json
import os
import re
import subprocess
import sys
import tempfile
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, urlparse


ROOT = Path(__file__).resolve().parent
EXTRACTOR = ROOT / "extract_branch.py"
REQUIREMENTS = ROOT / "requirements.txt"
LABEL_OUTPUT_DIR = Path(os.environ.get("CORONARY_LABEL_OUTPUT_DIR", str(ROOT / "labels")))


def ensure_dependencies():
    """Install the extractor's Python dependencies on first launch."""
    required = ("numpy", "scipy", "skimage", "nibabel", "openpyxl")
    missing = [name for name in required if importlib.util.find_spec(name) is None]
    if not missing:
        return
    print("Installing required Python packages: " + ", ".join(missing), flush=True)
    command = [sys.executable, "-m", "pip", "install", "-r", str(REQUIREMENTS)]
    result = subprocess.run(command, cwd=ROOT, text=True, capture_output=True)
    if result.returncode != 0:
        print(result.stdout, flush=True)
        print(result.stderr, flush=True)
        raise RuntimeError(
            "Could not install Python dependencies automatically. "
            "Check server.log for the pip error."
        )


class Handler(SimpleHTTPRequestHandler):
    def end_headers(self):
        self.send_header("Access-Control-Allow-Origin", "*")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
        super().end_headers()

    def do_OPTIONS(self):
        self.send_response(204)
        self.end_headers()

    def do_GET(self):
        parsed = urlparse(self.path)
        if parsed.path == "/api/load-folder":
            self.load_folder(parsed)
            return
        if parsed.path == "/api/export-excel":
            self.export_excel(parsed)
            return
        super().do_GET()

    @staticmethod
    def _safe_component(value):
        return (
            isinstance(value, str)
            and 1 <= len(value) <= 120
            and "\0" not in value
            and not any(char in value for char in "/\\")
            and value not in (".", "..")
        )

    def _progress_path(self, folder):
        return LABEL_OUTPUT_DIR / f"{folder}.json"

    @staticmethod
    def _clean_label(value):
        if isinstance(value, bool):
            return value
        if isinstance(value, dict):
            return {
                "present": str(value.get("present") or "")[:20],
                "score": str(value.get("score") or "")[:20],
                "plaque": str(value.get("plaque") or "")[:20],
            }
        return None

    def load_folder(self, parsed):
        folder = (parse_qs(parsed.query).get("folder") or [""])[0]
        if not self._safe_component(folder):
            self._json({"error": "Invalid folder name."}, 400)
            return
        path = self._progress_path(folder)
        if not path.exists():
            self._json({"folder": folder, "masks": {}}, 200)
            return
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            self._json({"error": f"Could not read saved progress: {exc}"}, 500)
            return
        masks = {}
        for entry in data.get("masks", []):
            if isinstance(entry, dict) and self._safe_component(entry.get("name")):
                masks[entry["name"]] = entry
        self._json({"folder": folder, "masks": masks})

    def save_progress(self):
        try:
            length = int(self.headers.get("Content-Length", "0"))
            payload = json.loads(self.rfile.read(length).decode("utf-8"))
        except (ValueError, json.JSONDecodeError) as exc:
            self._json({"error": f"Invalid JSON payload: {exc}"}, 400)
            return
        folder = payload.get("folder")
        masks = payload.get("masks")
        if not self._safe_component(folder):
            self._json({"error": "Invalid folder name."}, 400)
            return
        if not isinstance(masks, list) or len(masks) > 2000:
            self._json({"error": "Expected a masks array."}, 400)
            return
        cleaned = []
        for item in masks:
            if not isinstance(item, dict) or not self._safe_component(item.get("name")):
                continue
            labels = item.get("labels")
            clean_labels = {}
            if isinstance(labels, dict) and len(labels) <= 1000:
                for key, value in labels.items():
                    if not isinstance(key, str) or len(key) > 200:
                        continue
                    label = self._clean_label(value)
                    if label is not None:
                        clean_labels[key] = label
            cleaned.append({
                "name": item["name"],
                "order": int(item.get("order") or 0),
                "finished": bool(item.get("finished")),
                "finishedAt": str(item.get("finishedAt") or "")[:40],
                "notes": str(item.get("notes") or "")[:5000],
                "branchTotal": int(item.get("branchTotal") or 0),
                "labels": clean_labels,
            })
        try:
            LABEL_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
            path = self._progress_path(folder)
            temporary = path.with_suffix(".json.tmp")
            with temporary.open("w", encoding="utf-8") as handle:
                json.dump({"folder": folder, "masks": cleaned}, handle)
                handle.flush()
                os.fsync(handle.fileno())
            os.replace(temporary, path)
            self._json({"saved": True, "masks": len(cleaned)})
        except OSError as exc:
            self._json({"error": str(exc)}, 500)

    SCORE_RANK = {"1": 1, "2": 2, "3": 3, "4A": 4, "5": 5, "N": 6}
    PLAQUE_RANK = {"P1": 1, "P2": 2, "P3": 3, "P4": 4}

    @classmethod
    def _branch_states(cls, entry):
        labels = entry.get("labels") or {}
        states = []
        for key, value in labels.items():
            if "-" not in key:
                continue
            index_text, _, branch_name = key.partition("-")
            try:
                index = int(index_text)
            except ValueError:
                continue
            if isinstance(value, bool):
                state = {"present": "Yes" if value else "No", "score": "", "plaque": ""}
            elif isinstance(value, dict):
                state = {
                    "present": str(value.get("present") or ""),
                    "score": str(value.get("score") or ""),
                    "plaque": str(value.get("plaque") or ""),
                }
            else:
                continue
            states.append((index, branch_name, state))
        states.sort(key=lambda item: item[0])
        return states

    @classmethod
    def _excel_row(cls, entry):
        states = cls._branch_states(entry)
        total = int(entry.get("branchTotal") or 0)
        if not total and states:
            total = states[-1][0] + 1
        presents = [state["present"] for _, _, state in states]
        labeled_count = sum(1 for present in presents if present)
        if not states:
            overall = ""
        elif "Yes" in presents:
            overall = "Yes"
        elif "Unsure" in presents:
            overall = "Unsure"
        elif labeled_count and labeled_count < max(total, len(states)):
            overall = "Partial"
        elif labeled_count:
            overall = "No"
        else:
            overall = ""
        scores = [state["score"] for _, _, state in states if state["present"] == "Yes" and state["score"]]
        best_score = max(scores, key=lambda s: cls.SCORE_RANK.get(s, 0)) if scores else ""
        plaques = [state["plaque"] for _, _, state in states if state["present"] == "Yes" and state["plaque"]]
        best_plaque = max(plaques, key=lambda p: cls.PLAQUE_RANK.get(p, 0)) if plaques else ""
        if entry.get("finished"):
            status = "Finished"
        elif states:
            status = "In progress"
        else:
            status = "Pending"
        sides = {}
        details = []
        by_index = {index: (name, state) for index, name, state in states}
        for index in range(total):
            found = by_index.get(index)
            name = found[0] if found else ""
            state = found[1] if found else None
            if re.search(r"_l\.nii", name, re.I):
                side = "L"
            elif re.search(r"_r\.nii", name, re.I):
                side = "R"
            else:
                side = "B"
            sides[side] = sides.get(side, 0) + 1
            label = f"{side}{sides[side]}"
            if state is None or not state["present"]:
                details.append(f"{label}: Unlabeled")
                continue
            present = state["present"]
            if present == "Yes":
                extras = [f"score {state['score'] or '—'}"]
                if state["plaque"]:
                    extras.append(state["plaque"])
                details.append(f"{label}: Yes ({', '.join(extras)})")
            else:
                details.append(f"{label}: {present}")
        branches = f"{labeled_count}/{total}" if total else "—"
        notes = str(entry.get("notes") or "")
        return [
            entry.get("order") or 0,
            entry.get("name"),
            status,
            overall,
            best_score or "—",
            best_plaque or "—",
            branches,
            "; ".join(details) if details else "—",
            notes,
            entry.get("finishedAt") or "",
        ]

    def export_excel(self, parsed):
        folder = (parse_qs(parsed.query).get("folder") or [""])[0]
        if not self._safe_component(folder):
            self._json({"error": "Invalid folder name."}, 400)
            return
        path = self._progress_path(folder)
        if not path.exists():
            self._json({"error": "No saved cases for this folder yet."}, 404)
            return
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            self._json({"error": f"Could not read saved progress: {exc}"}, 500)
            return
        entries = [m for m in data.get("masks", []) if isinstance(m, dict) and self._safe_component(m.get("name"))]
        if not entries:
            self._json({"error": "No saved cases for this folder yet."}, 404)
            return
        entries.sort(key=lambda m: (int(m.get("order") or 0), m.get("name") or ""))
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
            from openpyxl.utils import get_column_letter
        except ImportError:
            self._json({"error": "openpyxl is not installed. Restart the server to install it."}, 500)
            return
        headers = ["#", "Mask file", "Status", "CAD present", "CAD-RADS score", "Plaque burden",
                   "Branches", "Branch detail", "Case notes", "Finished at"]
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Cases"
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for position, entry in enumerate(entries, start=1):
            row = self._excel_row(entry)
            sheet.append([position] + row[1:])
        widths = [6, 28, 13, 13, 16, 16, 11, 64, 30, 22]
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"
        buffer = io.BytesIO()
        workbook.save(buffer)
        raw = buffer.getvalue()
        safe_name = re.sub(r"[^\w.-]+", "_", folder) or "export"
        self.send_response(200)
        self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.send_header("Content-Length", str(len(raw)))
        self.send_header("Content-Disposition", f'attachment; filename="{safe_name}.xlsx"')
        self.end_headers()
        self.wfile.write(raw)

    def do_POST(self):
        route = urlparse(self.path).path
        if route == "/api/save-labels":
            self.save_labels()
            return
        if route == "/api/save-progress":
            self.save_progress()
            return
        if route != "/api/extract-branches":
            self.send_error(404, "Unknown API endpoint")
            return
        length = int(self.headers.get("Content-Length", "0"))
        body = self.rfile.read(length)
        content_type = self.headers.get("Content-Type", "")
        boundary = content_type.split("boundary=", 1)[-1].strip().strip('"')
        if not boundary:
            self._json({"error": "Expected multipart/form-data with a mask field."}, 400)
            return
        marker = ("--" + boundary).encode()
        upload = None
        filename = "coronary_mask.nii.gz"
        for part in body.split(marker):
            if b"name=\"mask\"" not in part:
                continue
            head, sep, data = part.partition(b"\r\n\r\n")
            if not sep:
                continue
            for line in head.decode("utf-8", "ignore").split("\r\n"):
                if "filename=" in line:
                    filename = line.split("filename=", 1)[1].strip().strip('"')
            upload = data[:-2] if data.endswith(b"\r\n") else data
            break
        if not upload:
            self._json({"error": "No mask file was uploaded."}, 400)
            return

        try:
            with tempfile.TemporaryDirectory(prefix="coronary-labeler-") as temp:
                work = Path(temp)
                suffix = ".nii.gz" if filename.lower().endswith(".nii.gz") else Path(filename).suffix
                mask_path = work / ("uploaded_mask" + suffix)
                out_dir = work / "branches"
                mask_path.write_bytes(upload)
                interpreter = ["py", "-3"] if os.name == "nt" else [sys.executable]
                command = interpreter + [str(EXTRACTOR), str(mask_path), "--out-dir", str(out_dir)]
                result = subprocess.run(command, cwd=ROOT, text=True, capture_output=True, timeout=300)
                if result.returncode != 0:
                    message = (result.stderr or result.stdout or "Branch extraction failed").strip()
                    print(f"Branch extraction failed: {message}", flush=True)
                    self._json({"error": message}, 422)
                    return
                files = sorted(out_dir.glob("*/branch_*_[lru].nii.gz"), key=self._branch_sort_key)
                if not files:
                    self._json({"error": "The extractor did not produce any branch masks."}, 422)
                    return
                branches = []
                for path in files:
                    name = path.name
                    side = "left" if name.rsplit("_", 1)[-1].startswith("l") else "right" if name.rsplit("_", 1)[-1].startswith("r") else "unsplit"
                    branches.append({
                        "name": name,
                        "side": side,
                        "data": base64.b64encode(path.read_bytes()).decode("ascii"),
                    })
                self._json({"branches": branches, "log": result.stdout[-4000:]})
        except subprocess.TimeoutExpired:
            print("Branch extraction timed out after 5 minutes.", flush=True)
            self._json({"error": "Branch extraction timed out after 5 minutes."}, 504)
        except Exception as exc:
            print(f"Branch extraction server error: {exc}", flush=True)
            self._json({"error": str(exc)}, 500)

    def save_labels(self):
        try:
            length = int(self.headers.get("Content-Length", "0"))
            payload = json.loads(self.rfile.read(length).decode("utf-8"))
            case_id = str(payload.get("case_id", ""))
            import re
            if not re.fullmatch(r"CVD\d{1,40}", case_id):
                self._json({"error": "Invalid anonymised case ID."}, 400)
                return
            rows = payload.get("rows")
            if not isinstance(rows, list):
                self._json({"error": "Expected a rows array."}, 400)
                return
            fields = ["case_id", "branch_number", "cadrads_present", "stenosis", "cadrads_score", "plaque_burden", "case_notes", "labeled_at"]
            LABEL_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
            destination = LABEL_OUTPUT_DIR / f"{case_id}_labels.csv"
            temporary = destination.with_suffix(destination.suffix + ".tmp")
            with temporary.open("w", newline="", encoding="utf-8") as handle:
                writer = csv.DictWriter(handle, fieldnames=fields, extrasaction="ignore")
                writer.writeheader()
                for row in rows:
                    writer.writerow({field: row.get(field, "") for field in fields})
                handle.flush()
                os.fsync(handle.fileno())
            os.replace(temporary, destination)
            self._json({"saved": True, "path": str(destination), "rows": len(rows)})
        except Exception as exc:
            self._json({"error": str(exc)}, 500)

    @staticmethod
    def _branch_sort_key(path):
        name = path.name
        side = name.rsplit("_", 1)[-1][0]
        try:
            number = int(name.split("_", 2)[1])
        except (IndexError, ValueError):
            number = 999999
        return ("lru".index(side) if side in "lru" else 9, number)

    def _json(self, payload, status=200):
        raw = json.dumps(payload).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(raw)))
        self.end_headers()
        self.wfile.write(raw)


if __name__ == "__main__":
    os.chdir(ROOT)
    ensure_dependencies()
    port = int(os.environ.get("CORONARY_LABEL_PORT", "8765"))
    server = ThreadingHTTPServer(("127.0.0.1", port), Handler)
    print(f"Coronary branch labeler: http://127.0.0.1:{port}")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\nStopped.")
    finally:
        server.server_close()
